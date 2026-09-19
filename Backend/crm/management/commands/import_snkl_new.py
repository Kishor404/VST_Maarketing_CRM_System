import os
import re
from datetime import datetime, date

from django.core.management.base import BaseCommand
from django.contrib.auth import get_user_model

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

# IMPORTANT:
# Change this if Card is inside another app.
from crm.models import Card


User = get_user_model()


class Command(BaseCommand):
    help = "Import customers and cards from Excel"

    REGION = "sankarankovil"

    # =========================================================
    # ARGUMENTS
    # =========================================================

    def add_arguments(self, parser):

        parser.add_argument(
            "input_file",
            type=str,
            help="Path to input Excel file",
        )

        parser.add_argument(
            "--output",
            type=str,
            default=None,
            help="Output Excel file path",
        )

    # =========================================================
    # MAIN
    # =========================================================

    def handle(self, *args, **options):

        input_file = options["input_file"]
        output_file = options["output"]

        # -----------------------------------------------------
        # Check input file
        # -----------------------------------------------------

        if not os.path.exists(input_file):

            self.stdout.write(
                self.style.ERROR(
                    f"Input file not found: {input_file}"
                )
            )

            return

        if not input_file.lower().endswith(".xlsx"):

            self.stdout.write(
                self.style.ERROR(
                    "Only .xlsx files are supported."
                )
            )

            return

        # -----------------------------------------------------
        # Generate output filename
        # -----------------------------------------------------

        if not output_file:

            directory = os.path.dirname(input_file)

            filename = os.path.basename(input_file)

            filename_without_extension = os.path.splitext(
                filename
            )[0]

            output_file = os.path.join(
                directory,
                f"{filename_without_extension}_result.xlsx"
            )

        # =====================================================
        # LOAD EXCEL
        # =====================================================

        self.stdout.write("")

        self.stdout.write(
            self.style.WARNING(
                f"Reading Excel: {input_file}"
            )
        )

        workbook = load_workbook(
            input_file
        )

        worksheet = workbook.active

        # =====================================================
        # READ HEADERS
        # =====================================================

        headers = {}

        for column in range(
            1,
            worksheet.max_column + 1
        ):

            value = worksheet.cell(
                row=1,
                column=column
            ).value

            if value is not None:

                header = str(value).strip().lower()

                headers[header] = column

        # =====================================================
        # REQUIRED COLUMNS
        # =====================================================

        required_columns = [
            "lid",
            "name",
            "address",
            "city",
            "phone",
            "model",
            "date_of_installation",
            "warranty_start_date",
            "warranty_end_date",
        ]

        missing_columns = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing_columns:

            self.stdout.write(
                self.style.ERROR(
                    "Missing required columns: "
                    + ", ".join(missing_columns)
                )
            )

            return

        # =====================================================
        # RESULT COLUMNS
        # =====================================================

        original_max_column = worksheet.max_column

        status_column = original_max_column + 1

        customer_code_column = original_max_column + 2

        user_action_column = original_max_column + 3

        error_column = original_max_column + 4

        worksheet.cell(
            row=1,
            column=status_column,
            value="status"
        )

        worksheet.cell(
            row=1,
            column=customer_code_column,
            value="customer_code"
        )

        worksheet.cell(
            row=1,
            column=user_action_column,
            value="user_action"
        )

        worksheet.cell(
            row=1,
            column=error_column,
            value="error"
        )

        # -----------------------------------------------------
        # Style headers
        # -----------------------------------------------------

        for column in [
            status_column,
            customer_code_column,
            user_action_column,
            error_column,
        ]:

            cell = worksheet.cell(
                row=1,
                column=column
            )

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        # =====================================================
        # COUNTERS
        # =====================================================

        success_count = 0

        skipped_count = 0

        user_failed_count = 0

        card_failed_count = 0

        new_user_count = 0

        existing_user_count = 0

        # =====================================================
        # PROCESS EACH ROW
        # =====================================================

        for row_number in range(
            2,
            worksheet.max_row + 1
        ):

            status_text = ""

            customer_code = ""

            user_action = ""

            error_message = ""

            # =================================================
            # READ DATA
            # =================================================

            lid = self.get_value(
                worksheet,
                row_number,
                headers,
                "lid"
            )

            excel_name = self.get_value(
                worksheet,
                row_number,
                headers,
                "name"
            )

            address = self.clean_text(
                self.get_value(
                    worksheet,
                    row_number,
                    headers,
                    "address"
                )
            )

            city = self.clean_text(
                self.get_value(
                    worksheet,
                    row_number,
                    headers,
                    "city"
                )
            )

            raw_phone = self.get_value(
                worksheet,
                row_number,
                headers,
                "phone"
            )

            model_name = self.clean_text(
                self.get_value(
                    worksheet,
                    row_number,
                    headers,
                    "model"
                )
            )

            raw_installation_date = self.get_value(
                worksheet,
                row_number,
                headers,
                "date_of_installation"
            )

            raw_warranty_start = self.get_value(
                worksheet,
                row_number,
                headers,
                "warranty_start_date"
            )

            raw_warranty_end = self.get_value(
                worksheet,
                row_number,
                headers,
                "warranty_end_date"
            )

            # =================================================
            # EMPTY ROW
            # =================================================

            if all(
                value in (None, "")
                for value in [
                    lid,
                    excel_name,
                    raw_phone,
                    model_name,
                    raw_installation_date,
                ]
            ):

                continue

            # =================================================
            # VALIDATE LID
            # =================================================

            if lid in (None, ""):

                status_text = "User Creation Failed"

                error_message = "lid is empty"

                user_failed_count += 1

                self.write_result(
                    worksheet,
                    row_number,
                    status_column,
                    customer_code_column,
                    user_action_column,
                    error_column,
                    status_text,
                    customer_code,
                    user_action,
                    error_message,
                )

                continue

            # =================================================
            # VALIDATE NAME
            # =================================================

            if excel_name in (None, ""):

                status_text = "User Creation Failed"

                error_message = "name is empty"

                user_failed_count += 1

                self.write_result(
                    worksheet,
                    row_number,
                    status_column,
                    customer_code_column,
                    user_action_column,
                    error_column,
                    status_text,
                    customer_code,
                    user_action,
                    error_message,
                )

                continue

            # =================================================
            # CREATE USER NAME
            #
            # Example:
            #
            # lid = 2282
            # name = RAJA MOHAMAD
            #
            # => 2282 RAJA MOHAMAD
            # =================================================

            lid_string = self.clean_lid(
                lid
            )

            customer_name = (
                f"{lid_string} "
                f"{str(excel_name).strip()}"
            )

            # =================================================
            # PHONE
            # =================================================

            phone = self.clean_phone(
                raw_phone
            )

            # =================================================
            # PARSE INSTALLATION DATE
            # =================================================

            try:

                installation_date = self.parse_date(
                    raw_installation_date
                )

                if not installation_date:

                    raise ValueError(
                        "date_of_installation is empty"
                    )

                warranty_start_date = self.parse_date(
                    raw_warranty_start
                )

                warranty_end_date = self.parse_date(
                    raw_warranty_end
                )

            except Exception as exc:

                status_text = "User Creation Failed"

                error_message = (
                    f"Invalid date: {str(exc)}"
                )

                user_failed_count += 1

                self.write_result(
                    worksheet,
                    row_number,
                    status_column,
                    customer_code_column,
                    user_action_column,
                    error_column,
                    status_text,
                    customer_code,
                    user_action,
                    error_message,
                )

                self.stdout.write(
                    self.style.ERROR(
                        f"Row {row_number}: "
                        f"{error_message}"
                    )
                )

                continue

            # =================================================
            # CREATE OR FIND USER
            # =================================================

            user = None

            try:

                # =================================================
                # PHONE EXISTS
                # =================================================

                if phone:

                    normalized_phone = (
                        User.objects.normalize_phone(
                            phone
                        )
                    )

                    # -------------------------------------------------
                    # FIND EXISTING USER
                    # -------------------------------------------------

                    user = User.objects.filter(
                        phone=normalized_phone
                    ).first()

                    # =================================================
                    # EXISTING USER FOUND
                    # =================================================

                    if user:

                        customer_code = (
                            user.customer_code
                        )

                        user_action = (
                            "Existing User"
                        )

                        existing_user_count += 1

                        self.stdout.write(
                            self.style.WARNING(
                                f"Row {row_number}: "
                                f"Existing user -> "
                                f"{user.name} "
                                f"({user.customer_code})"
                            )
                        )

                    # =================================================
                    # USER DOES NOT EXIST
                    # =================================================

                    else:

                        user = User.objects.create_user(
                            phone=normalized_phone,
                            name=customer_name,
                            address=address,
                            city=city,
                            region=self.REGION,
                            role="customer",
                            is_industrial=False,
                            is_active=True,
                        )

                        customer_code = (
                            user.customer_code
                        )

                        user_action = "Created"

                        new_user_count += 1

                        self.stdout.write(
                            self.style.SUCCESS(
                                f"Row {row_number}: "
                                f"New user -> "
                                f"{user.name} "
                                f"({user.customer_code})"
                            )
                        )

                # =================================================
                # PHONE BLANK
                # =================================================

                else:

                    user = User.objects.create_user(
                        phone=None,
                        name=customer_name,
                        address=address,
                        city=city,
                        region=self.REGION,
                        role="customer",
                        is_industrial=False,
                        is_active=True,
                    )

                    customer_code = (
                        user.customer_code
                    )

                    user_action = "Created"

                    new_user_count += 1

                    self.stdout.write(
                        self.style.SUCCESS(
                            f"Row {row_number}: "
                            f"New user without phone -> "
                            f"{user.name} "
                            f"({user.customer_code})"
                        )
                    )

            # =================================================
            # USER CREATION FAILED
            # =================================================

            except Exception as exc:

                status_text = "User Creation Failed"

                error_message = str(exc)

                user_failed_count += 1

                self.write_result(
                    worksheet,
                    row_number,
                    status_column,
                    customer_code_column,
                    user_action_column,
                    error_column,
                    status_text,
                    customer_code,
                    user_action,
                    error_message,
                )

                self.stdout.write(
                    self.style.ERROR(
                        f"Row {row_number}: "
                        f"User creation failed -> "
                        f"{error_message}"
                    )
                )

                continue

            # =================================================
            # IMPORTANT:
            #
            # USER EXISTS
            # ↓
            # CHECK INSTALLATION DATE
            # =================================================

            try:

                existing_card = Card.objects.filter(
                    customer=user,
                    date_of_installation=installation_date
                ).first()

            except Exception as exc:

                status_text = (
                    "Card Creation failed but user created"
                )

                error_message = (
                    f"Unable to check existing card: "
                    f"{str(exc)}"
                )

                card_failed_count += 1

                self.write_result(
                    worksheet,
                    row_number,
                    status_column,
                    customer_code_column,
                    user_action_column,
                    error_column,
                    status_text,
                    customer_code,
                    user_action,
                    error_message,
                )

                continue

            # =================================================
            # SAME INSTALLATION DATE FOUND
            #
            # SKIP CARD CREATION
            # =================================================

            if existing_card:

                status_text = (
                    "Skipped - Card Already Exists"
                )

                user_action = (
                    f"{user_action}; "
                    f"Existing Card #{existing_card.id}"
                )

                skipped_count += 1

                self.stdout.write(
                    self.style.WARNING(
                        f"Row {row_number}: "
                        f"SKIPPED -> "
                        f"Card #{existing_card.id} "
                        f"already exists for "
                        f"{user.name} "
                        f"with installation date "
                        f"{installation_date}"
                    )
                )

                self.write_result(
                    worksheet,
                    row_number,
                    status_column,
                    customer_code_column,
                    user_action_column,
                    error_column,
                    status_text,
                    customer_code,
                    user_action,
                    "",
                )

                continue

            # =================================================
            # NO CARD WITH SAME INSTALLATION DATE
            #
            # CREATE NEW CARD
            # =================================================

            try:

                card = Card.objects.create(

                    model=model_name,

                    customer=user,

                    # Always use actual user's name
                    customer_name=user.name,

                    card_type="normal",

                    region=self.REGION,

                    address=address,

                    city=city,

                    postal_code="",

                    date_of_installation=installation_date,

                    # IMPORTANT:
                    # Use dates from CURRENT Excel
                    warranty_start_date=warranty_start_date,

                    warranty_end_date=warranty_end_date,

                    amc_start_date=None,

                    amc_end_date=None,
                )

                status_text = "Success"

                success_count += 1

                self.stdout.write(
                    self.style.SUCCESS(
                        f"Row {row_number}: "
                        f"Card #{card.id} created -> "
                        f"{user.name}"
                    )
                )

            # =================================================
            # CARD CREATION FAILED
            # =================================================

            except Exception as exc:

                status_text = (
                    "Card Creation failed but user created"
                )

                error_message = str(exc)

                card_failed_count += 1

                self.stdout.write(
                    self.style.ERROR(
                        f"Row {row_number}: "
                        f"Card creation failed -> "
                        f"{error_message}"
                    )
                )

            # =================================================
            # WRITE RESULT
            # =================================================

            self.write_result(
                worksheet,
                row_number,
                status_column,
                customer_code_column,
                user_action_column,
                error_column,
                status_text,
                customer_code,
                user_action,
                error_message,
            )

        # =====================================================
        # COLUMN WIDTH
        # =====================================================

        worksheet.column_dimensions[
            self.column_letter(status_column)
        ].width = 40

        worksheet.column_dimensions[
            self.column_letter(customer_code_column)
        ].width = 20

        worksheet.column_dimensions[
            self.column_letter(user_action_column)
        ].width = 35

        worksheet.column_dimensions[
            self.column_letter(error_column)
        ].width = 60

        # =====================================================
        # SAVE
        # =====================================================

        try:

            workbook.save(
                output_file
            )

        except PermissionError:

            self.stdout.write(
                self.style.ERROR(
                    "Could not save the result file. "
                    "Make sure the Excel file is not open."
                )
            )

            return

        # =====================================================
        # SUMMARY
        # =====================================================

        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                "========================================"
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                "IMPORT COMPLETED"
            )
        )

        self.stdout.write(
            self.style.SUCCESS(
                "========================================"
            )
        )

        self.stdout.write(
            f"New users created      : {new_user_count}"
        )

        self.stdout.write(
            f"Existing users reused  : {existing_user_count}"
        )

        self.stdout.write(
            f"Cards created          : {success_count}"
        )

        self.stdout.write(
            f"Cards skipped          : {skipped_count}"
        )

        self.stdout.write(
            f"User creation failed   : {user_failed_count}"
        )

        self.stdout.write(
            f"Card creation failed   : {card_failed_count}"
        )

        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                f"Result file: {output_file}"
            )
        )

    # =========================================================
    # GET VALUE
    # =========================================================

    @staticmethod
    def get_value(
        worksheet,
        row,
        headers,
        field_name
    ):

        column = headers[field_name]

        return worksheet.cell(
            row=row,
            column=column
        ).value

    # =========================================================
    # CLEAN TEXT
    # =========================================================

    @staticmethod
    def clean_text(value):

        if value is None:
            return ""

        return str(value).strip()

    # =========================================================
    # CLEAN LID
    # =========================================================

    @staticmethod
    def clean_lid(value):

        if value is None:
            return ""

        if isinstance(value, float):

            if value.is_integer():

                return str(
                    int(value)
                )

        return str(value).strip()

    # =========================================================
    # CLEAN PHONE
    # =========================================================

    @staticmethod
    def clean_phone(value):

        if value is None:
            return None

        # -----------------------------------------------------
        # Excel integer
        # -----------------------------------------------------

        if isinstance(value, int):

            value = str(value)

        # -----------------------------------------------------
        # Excel float
        # -----------------------------------------------------

        elif isinstance(value, float):

            if value.is_integer():

                value = str(
                    int(value)
                )

            else:

                value = str(value)

        # -----------------------------------------------------
        # String
        # -----------------------------------------------------

        else:

            value = str(value).strip()

        if not value:

            return None

        # -----------------------------------------------------
        # Remove .0 from Excel
        # -----------------------------------------------------

        if value.endswith(".0"):

            value = value[:-2]

        # -----------------------------------------------------
        # Keep + if international number
        # -----------------------------------------------------

        if value.startswith("+"):

            value = (
                "+"
                + re.sub(
                    r"\D",
                    "",
                    value[1:]
                )
            )

        else:

            value = re.sub(
                r"\D",
                "",
                value
            )

        if not value:

            return None

        return value

    # =========================================================
    # PARSE DATE
    # =========================================================

    @staticmethod
    def parse_date(value):

        if value is None or value == "":

            return None

        # -----------------------------------------------------
        # Excel datetime
        # -----------------------------------------------------

        if isinstance(value, datetime):

            return value.date()

        # -----------------------------------------------------
        # Python date
        # -----------------------------------------------------

        if isinstance(value, date):

            return value

        # -----------------------------------------------------
        # String
        # -----------------------------------------------------

        value = str(value).strip()

        formats = [
            "%Y-%m-%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%Y/%m/%d",
            "%d-%m-%y",
            "%d/%m/%y",
        ]

        for fmt in formats:

            try:

                return datetime.strptime(
                    value,
                    fmt
                ).date()

            except ValueError:

                continue

        raise ValueError(
            f"Unsupported date format: {value}"
        )

    # =========================================================
    # WRITE RESULT
    # =========================================================

    @staticmethod
    def write_result(
        worksheet,
        row_number,
        status_column,
        customer_code_column,
        user_action_column,
        error_column,
        status_text,
        customer_code,
        user_action,
        error_message,
    ):

        worksheet.cell(
            row=row_number,
            column=status_column,
            value=status_text
        )

        worksheet.cell(
            row=row_number,
            column=customer_code_column,
            value=customer_code
        )

        worksheet.cell(
            row=row_number,
            column=user_action_column,
            value=user_action
        )

        worksheet.cell(
            row=row_number,
            column=error_column,
            value=error_message
        )

    # =========================================================
    # COLUMN NUMBER -> LETTER
    # =========================================================

    @staticmethod
    def column_letter(column_number):

        result = ""

        while column_number:

            column_number, remainder = divmod(
                column_number - 1,
                26
            )

            result = (
                chr(65 + remainder)
                + result
            )

        return result